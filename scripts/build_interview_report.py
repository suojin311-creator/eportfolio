"""Rebuild the interview attachment as a work report from supplied interview evidence."""
from pathlib import Path
from xml.sax.saxutils import escape
from fontTools.ttLib import TTFont
from fontTools.varLib.instancer import instantiateVariableFont
from fontTools import subset
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor,white
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont as RLFont
from docx import Document
from docx.shared import Pt,Cm
import fitz
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'public/documents'
SUMMARY=[
'受访者认可部分育儿 AI 建议，希望得到可以直接采取的做法。',
'额外准备和互动步骤会占用亲子时间，建议采用起来是否省事，需要作为产品设计重点。',
'家庭已有语音定时习惯。智能玩偶需在主题互动和陪伴方式上提供额外价值。',
'家长对能力标签持谨慎态度，产品应区分一次练习记录与长期能力判断。']
FINDINGS=[
('需要具体做法','“我更需要的是你给我出一个解决方案，我应该怎么做？”','讨论情景模拟时，受访者认为已有育儿 AI 能提供方法；希望建议直接回应具体问题。'),
('执行受时间限制','“我未必有时间，他也未必有时间来做这件事情”','讨论手工、卡片等活动时，她认可部分建议，但会结合已有活动调整执行方式。'),
('已经使用语音定时','“我会直接让小爱同学定时”','家长先与孩子商量动画结束、洗澡等时间，再使用智能音箱定时。'),
('不希望给孩子贴标签','“我也拒绝他给我孩子打标签”','该意见涉及部分测评功能，应与她对其他育儿 AI 功能的认可一起理解。')]
RECOMMENDATIONS=[
('先提供一项可直接采用的活动','将建议整理为主题任务模板，默认配置可直接确认，调整提醒作为可选操作。','家长能否理解并采用，是否仍需额外构思或说明。'),
('活动尽量接入家庭已有安排','首轮用刷牙等日常活动试用，避免同时叠加多项准备要求。','家长愿意安排在哪个时点，执行与确认是否增加负担。'),
('单独验证智能玩偶的参与价值','在已有定时工具之外，观察主题提问和孩子主动触发是否有吸引力。','孩子是否主动发起，对谈能否围绕主题继续。'),
('反馈只说明记录得到的行为','练习记录与现实行为分开；实际完成由家长核验，避免从一次表现推断能力。','家长能否准确理解记录，是否认可确认方式。')]
LIMITS=[
'本次为 1 户老用户家庭。访谈主要讨论近 11 岁孩子，家庭另有 6 岁儿童；需补充 3–6 岁家庭观察。',
'本份访谈没有直接验证成长目标设置、NFC 卡片、套餐付费意愿或留存改善。',
'成长目标设置困难与智能玩偶对话表现等其他线索来自项目回忆，应另行核实。']
ACTIONS=[
'补充目标年龄家庭，了解日常刷牙安排及现有提醒方式。',
'用一项预填任务做家长试用，记录是否理解、是否采用及需要修改的位置。',
'单独观察孩子对主题提问的参与，再结合家庭实际执行判断是否继续迭代。']
meta='关联项目：小乐好习惯养成方案    项目周期：2026.03—2026.06'
sample='访谈对象：蛋蛋妈妈｜老用户、二孩家庭｜依据提供的访谈转写整理'
alltext='老用户访谈总结及产品建议结论摘要主要发现产品建议需要补充验证下一步建议发现原文依据背景与解释建议方向建议内容观察重点'+meta+sample+''.join(SUMMARY+LIMITS+ACTIONS)+''.join(''.join(r) for r in FINDINGS+RECOMMENDATIONS)+''.join(chr(i) for i in range(32,127))
font_src=Path('/workspace/research/fonts/NotoSansSC.ttf')
if not font_src.exists():raise SystemExit('NotoSansSC.ttf must be available under /workspace/research/fonts')
for weight,name in [(400,'ReportSans'),(600,'ReportSansBold')]:
 f=TTFont(font_src);instantiateVariableFont(f,{'wght':weight},inplace=True);o=subset.Options();o.layout_features=['*'];w=subset.Subsetter(options=o);w.populate(text=alltext+'访谈总结重构整理第页');w.subset(f);path=Path('/workspace/scratch')/(name+'.ttf');f.save(path);pdfmetrics.registerFont(RLFont(name,str(path)))
styles={
 'body':ParagraphStyle('body',fontName='ReportSans',fontSize=10.5,leading=17,wordWrap='CJK',textColor=HexColor('#28352f'),spaceAfter=7),
 'small':ParagraphStyle('small',fontName='ReportSans',fontSize=9,leading=15,wordWrap='CJK',textColor=HexColor('#68756c'),spaceAfter=5),
 'h1':ParagraphStyle('h1',fontName='ReportSansBold',fontSize=21,leading=30,spaceAfter=13,textColor=HexColor('#223d2c')),
 'h2':ParagraphStyle('h2',fontName='ReportSansBold',fontSize=13,leading=21,spaceBefore=13,spaceAfter=9,textColor=HexColor('#2c5138')),
 'cell':ParagraphStyle('cell',fontName='ReportSans',fontSize=9.3,leading=15.5,wordWrap='CJK',textColor=HexColor('#36423a')),
 'head':ParagraphStyle('head',fontName='ReportSansBold',fontSize=9.3,leading=15,textColor=HexColor('#2d4a35'))}
def P(t,style='body'):return Paragraph(escape(t),styles[style])
def table(head,rows,widths):
 data=[[P(x,'head') for x in head]]+[[P(x,'cell') for x in row] for row in rows]
 t=Table(data,colWidths=widths,hAlign='LEFT',repeatRows=1);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#edf3ed')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10),('LINEBELOW',(0,0),(-1,-1),.4,HexColor('#d8e2d9'))]));return t
story=[P('老用户访谈总结及产品建议','h1'),P(meta,'small'),P(sample,'small'),P('一、结论摘要','h2')]
for i,t in enumerate(SUMMARY,1):story.append(P(str(i)+'. '+t))
story.extend([P('二、主要发现','h2'),table(['发现','原文依据','背景与解释'],FINDINGS,[86,200,213]),PageBreak(),P('三、产品建议','h2'),table(['建议方向','建议内容','观察重点'],RECOMMENDATIONS,[106,210,183]),P('四、需要补充验证','h2')])
for i,t in enumerate(LIMITS,1):story.append(P(str(i)+'. '+t))
story.append(P('五、下一步建议','h2'))
for i,t in enumerate(ACTIONS,1):story.append(P(str(i)+'. '+t))
def footer(canvas,doc):
 canvas.saveState();canvas.setStrokeColor(HexColor('#d8e2d9'));canvas.line(48,42,547,42);canvas.setFont('ReportSans',8);canvas.setFillColor(HexColor('#7a877b'));canvas.drawString(48,27,'访谈总结 · 重构整理');canvas.drawRightString(547,27,str(doc.page));canvas.restoreState()
path=OUT/'case01-interview.pdf';doc=SimpleDocTemplate(str(path),pagesize=(595.28,841.89),rightMargin=48,leftMargin=48,topMargin=44,bottomMargin=55,title='老用户访谈总结及产品建议',author='');doc.build(story,onFirstPage=footer,onLaterPages=footer)
d=fitz.open(path);assert len(d)==2,len(d);metadata=d.metadata;metadata['creationDate']='';metadata['modDate']='';d.set_metadata(metadata);temp=path.with_suffix('.new.pdf');d.save(temp,garbage=4,deflate=True);d.close();temp.replace(path)
# Editable working copy uses the same content and work-report structure.
doc=Document();section=doc.sections[0];section.top_margin=Cm(1.7);section.bottom_margin=Cm(1.7);section.left_margin=Cm(1.7);section.right_margin=Cm(1.7);normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10.5)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
doc.add_heading('老用户访谈总结及产品建议',0);doc.add_paragraph(meta);doc.add_paragraph(sample);doc.add_heading('一、结论摘要',1)
for t in SUMMARY:doc.add_paragraph(t,style='List Number')
def word_table(head,rows):
 t=doc.add_table(rows=1,cols=len(head));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,head):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=v
doc.add_heading('二、主要发现',1);word_table(['发现','原文依据','背景与解释'],FINDINGS);doc.add_page_break();doc.add_heading('三、产品建议',1);word_table(['建议方向','建议内容','观察重点'],RECOMMENDATIONS)
doc.add_heading('四、需要补充验证',1)
for t in LIMITS:doc.add_paragraph(t,style='List Number')
doc.add_heading('五、下一步建议',1)
for t in ACTIONS:doc.add_paragraph(t,style='List Number')
doc.sections[0].footer.paragraphs[0].text='访谈总结 · 重构整理';doc.core_properties.title='老用户访谈总结及产品建议';
for prop in ['dcterms:created','dcterms:modified']:
 node=doc.core_properties._element.find(qn(prop))
 if node is not None:doc.core_properties._element.remove(node)
doc.save(OUT/'case01-interview.docx')
print('Interview work report built: 2-page PDF + editable DOCX')
